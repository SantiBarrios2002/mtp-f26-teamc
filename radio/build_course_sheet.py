"""Rebuild the course AN5142 link-budget sheet in its own layout, filled with our 869 MHz baseline.

The course file (docs/AN5142-link-budget_corr_211022.xls, not in the repo) is an old .xls
whose formulas cannot be edited here without Excel or LibreOffice. This script recreates its
"Link Budget" sheet row for row (same labels, variables, units and formulas). The course's
"Ground Multipath" sheet is inlined below row 43, and a few clearly marked Team C rows are added.

Sheets:
  - "Course example 2.4 GHz": the course's own inputs; must reproduce its values (check).
  - "SRI 70 m", "MRM 260 m", "NM hop 100 m": RFM69HCW baseline at 869.525 MHz, 100 kbps.

Run (needs openpyxl):  python3 radio/build_course_sheet.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).with_name("AN5142_course_layout_869MHz.xlsx")

F_TITLE = Font(bold=True, size=13)
F_HEAD = Font(bold=True)
F_IN = Font(color="1F4E9E")             # blue = input
F_NOTE = Font(italic=True, color="666666")
FILL_ASM = PatternFill("solid", fgColor="FFF2CC")   # yellow = assumption to measure
FILL_ADD = PatternFill("solid", fgColor="E3EFF1")   # Team C additions

COURSE = dict(f0=2467.0, ppa=10.0, lmt=None, lct1=-0.57, lcab_t=None, lct2=None, gt=-15.0, d=230.0,
              l0=0.0, lmp=None, lobs=0.0, gr=-15.0, lcr1=-0.57, lcab_r=None, lcr2=None, sens=-114.0,
              nf=7.0, t0=290.0, k=1.38e-23, bw=0.012, tant=300.0, htx=1.0, hrx=1.0, a=1.0)
BASE = dict(f0=869.525, ppa=17.0, lmt=-0.5, lct1=0.0, lcab_t=None, lct2=None, gt=0.0, d=70.0,
            l0=-1.0, lmp=0.0, lobs=0.0, gr=0.0, lcr1=0.0, lcab_r=None, lcr2=None, sens=-100.0,
            nf=7.0, t0=290.0, k=1.38e-23, bw=0.2, tant=300.0, htx=0.75, hrx=0.75, a=1.0)
SHEETS = [
    ("Course example 2.4 GHz", COURSE,
     "Course inputs (Maxim AN5142 example as given in the course .xls). Values must match the original file."),
    ("SRI 70 m", dict(BASE, d=70.0),
     "Team C baseline: RFM69HCW 868, +17 dBm, 100 kbps GFSK, 869.525 MHz, antennas 0 dBi at 0.75 m, -1 dB container."),
    ("MRM 260 m", dict(BASE, d=260.0), "Team C baseline, MRM distance (~260 m per the rules)."),
    ("NM hop 100 m", dict(BASE, d=100.0, lobs=-10.0),
     "Team C baseline, one Network Mode hop of 100 m with -10 dB obstruction (no line of sight)."),
]

ASSUMPTIONS = {"gt", "gr", "l0", "sens", "htx", "hrx", "bw", "nf"}

# (row, label, variable, units, equation text, input key or None, formula/constant for non-inputs)
ROWS = [
    (3, "Frequency", "f0", "MHz", "", "f0", None),
    (4, "Speed of Light", "c", "m/s", "", None, 299792458),
    (5, "Wavelength", "λ", "m", "λ = c/f0", None, "=E4/(E3*1000000)"),
    (9, "PA Power", "PPA", "dBm", "", "ppa", None),
    (10, "TX Match Loss", "LMatchT", "dB", "", "lmt", None),
    (11, "TX source", "PTX", "dBm", "PTX = PPA + LMatchT", None, "=E9+E10"),
    (12, "TX connector loss", "LConT1", "dB", "(from Connector Loss sheet)", "lct1", None),
    (13, "TX cable loss", "LCabT", "dB", "(from Cable Loss sheet)", "lcab_t", None),
    (14, "TX connector loss (remote antenna)", "LConT2", "dB", "(from Connector Loss sheet)", "lct2", None),
    (15, "TX power", "PT", "dBm", "PT = PTX + (C&C Loss)", None, "=E11+E12+E13+E14"),
    (16, "TX antenna gain", "GT", "dBi", "", "gt", None),
    (17, "Effective (Isotropic) Radiated Power", "EIRP", "dBm", "EIRP = PT + GT", None, "=E15+E16"),
    (18, "Distance", "d", "m", "", "d", None),
    (19, "Channel Medium Loss Factor", "L0", "dB", "(from Medium Loss sheet)", "l0", None),
    (20, "Free Space Loss", "LFS", "dB", "LFS = (λ/4πd)²", None, "=20*LOG10(E5/(4*PI()*E18))"),
    (21, "Power at RX Antenna, Free Space Path", "PChanFS", "dB", "PChanFS = LFS + L0 + EIRP", None, "=E20+E19+E17"),
    (22, "Flat Earth Loss (Includes Ground Bounce)", "LFE", "dB", "(from Ground Multipath block, row 59)", None, "=E59"),
    (23, "Multipath Loss", "LMP", "dB", "", "lmp", None),
    (24, "Obstruction Loss", "LObs-Total", "dB", "", "lobs", None),
    (25, "Power at RX Antenna, Flat Earth Path", "PChanFE", "dB", "PChanFE = LFE + L0 + LMP + LObs + EIRP", None,
     "=E22+E19+E23+E24+E17"),
    (26, "RX antenna gain", "GR", "dBi", "", "gr", None),
    (27, "RX connector loss", "LConR1", "dB", "", "lcr1", None),
    (28, "RX cable loss", "LCabR", "dB", "", "lcab_r", None),
    (29, "RX connector loss (remote antenna)", "LConR2", "dB", "", "lcr2", None),
    (30, "RX power, Free Space Path", "PRFS", "dBm", "PRFS = PChanFS + GR + (C&C Loss)", None, "=E21+E26+E27+E28+E29"),
    (31, "RX power, Flat Earth Path", "PRFE", "dBm", "PRFE = PChanFE + GR + (C&C Loss)", None, "=E25+E26+E27+E28+E29"),
    (36, "RX Noise Figure", "NF", "dB", "", "nf", None),
    (37, "Operating Temperature", "T0", "K", "", "t0", None),
    (38, "Effective Noise Temperature", "Te", "K", "Te = T0(NF - 1)", None, "=E37*(10^(E36/10)-1)"),
    (39, "Boltzmann's constant", "k", "J/K", "", "k", None),
    (40, "Receive Bandwidth", "BWRX", "MHz", "", "bw", None),
    (41, "Antenna Temperature", "TAnt", "K", "", "tant", None),
    (42, "Noise Power (at RX)", "Pn", "dBm", "Pn = k(TAnt + Te)BWRX", None, "=10*LOG10(E39*(E41+E38)*E40*1000000)+30"),
    (43, "Signal to Noise Ratio", "SNRRX", "dB", "SNRRX = PRX/Pn  (course: sensitivity G31 over Pn)", None, "=G31-E42"),
    # Ground Multipath sheet of the course, inlined
    (48, "TX height", "hTX", "m", "", "htx", None),
    (49, "RX height", "hRX", "m", "", "hrx", None),
    (50, "Distance", "d", "m", "(from Link Budget, row 18)", None, "=E18"),
    (51, "Direct Path", "d1", "m", "d1 = (d² + (hTX - hRX)²)^½", None, "=SQRT(E50^2+(E48-E49)^2)"),
    (52, "Reflection Path", "d2", "m", "d2 = (d² + (hTX + hRX)²)^½", None, "=SQRT(E50^2+(E48+E49)^2)"),
    (53, "Reflection Path Δ", "Δd", "m", "Δd = d2 - d1", None, "=E52-E51"),
    (54, "Estimated Reflection Path Δ", "Δde", "m", "Δde = 2·hTX·hRX/d", None, "=2*E48*E49/E50"),
    (55, "Wavelength", "λ", "m", "(from Link Budget, row 5)", None, "=E5"),
    (56, "Normalized Ground Bounce Amplitude", "a", "none", "", "a", None),
    (57, "Loss From Ground Bounce Cancellation", "LGB", "dB", "LGB = 10log10(1 + a² - 2a|cos(2πΔd/λ)|)", None,
     "=10*LOG10(1+E56^2-2*E56*ABS(COS(2*PI()*E53/E55)))"),
    (58, "Free Space Loss", "LFS", "dB", "LFS = 10log10(λ/4πd)²", None, "=20*LOG10(E55/(4*PI()*E50))"),
    (59, "Flat Earth Loss", "LFE", "dB", "LFE = LGB + LFS", None, "=E57+E58"),
    # Team C additions
    (63, "Link margin, flat earth", "M", "dB", "M = PRFE - Sensitivity", None, "=E31-G31"),
    (64, "Link margin with AN5142's -20 dB multipath", "M20", "dB", "M20 = M - 20", None, "=E63-20"),
    (65, "SNR at PRFE", "SNR", "dB", "SNR = PRFE - Pn", None, "=E31-E42"),
    (66, "Effective Radiated Power", "ERP", "dBm", "ERP = EIRP - 2.15", None, "=E17-2.15"),
    (67, "Regulatory limit (869.40-869.65 MHz: 500 mW e.r.p.)", "Plim", "dBm", "CNAF UN-39", None, 27),
    (68, "Regulatory check", "", "", "ERP <= Plim", None, '=IF(E66<=E67,"OK","OVER LIMIT")'),
]


def build_sheet(ws, inputs: dict, note: str) -> None:
    ws["A1"] = "RF System (Link Budget) Calculations"
    ws["A1"].font = F_TITLE
    ws["G1"] = note
    ws["G1"].font = F_NOTE
    for r in (2, 8, 35, 47, 62):
        for c, h in enumerate(["Block", "Variable", "Units", "Equation", "Value"], start=1):
            ws.cell(row=r, column=c, value=h).font = F_HEAD
    ws["A2"] = "System Variables"
    ws["A35"] = "Receiver Sensitivity Calculations"
    ws["A46"] = "Ground Reflection Path Loss (course 'Ground Multipath' sheet, inlined here)"
    ws["A46"].font = F_HEAD
    ws["A61"] = "Added by Team C (not in the course sheet)"
    ws["A61"].font = F_HEAD

    is_course = inputs is COURSE
    for row, label, var, units, eq, key, f in ROWS:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=var)
        ws.cell(row=row, column=3, value=units)
        ws.cell(row=row, column=4, value=eq)
        cell = ws.cell(row=row, column=5)
        if key is not None:
            cell.value = inputs[key]
            cell.font = F_IN
            if key in ASSUMPTIONS and not is_course:
                cell.fill = FILL_ASM
        else:
            cell.value = f
        if row >= 63:
            for c in range(1, 6):
                ws.cell(row=row, column=c).fill = FILL_ADD
        cell.number_format = "0.00E+00" if key == "k" else "0.00##"

    ws["G31"] = inputs["sens"]
    ws["G31"].font = F_IN
    if not is_course:
        ws["G31"].fill = FILL_ASM
    ws["F32"] = "Sensitivity of Rx"
    ws["F32"].font = F_HEAD

    for col, width in {"A": 44, "B": 11, "C": 8, "D": 48, "E": 14, "G": 14}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="center")


def main() -> None:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    lines = [
        ("Course AN5142 link-budget sheet, rebuilt in its own layout with Team C's 869 MHz baseline", F_TITLE),
        ("Generated by radio/build_course_sheet.py. Same rows, labels and formulas as the course file", None),
        ("AN5142-link-budget_corr_211022.xls ('Link Budget' sheet); its 'Ground Multipath' sheet is inlined in rows 46-59.", None),
        ("Course correction kept: flat earth with no 1/2 factor and the |cos| worst-case ground bounce.", None),
        ("", None),
        ("Sheet 'Course example 2.4 GHz' re-enters the course's own inputs and reproduces its values (check).", None),
        ("Sheets 'SRI 70 m', 'MRM 260 m', 'NM hop 100 m': RFM69HCW 868 at +17 dBm, 100 kbps, 869.525 MHz.", None),
        ("", None),
        ("Blue = input. Yellow = assumption to replace with a measurement (antenna gain, container loss,", None),
        ("sensitivity at 100 kbps (datasheet stops at 38.4 kbps), heights, noise figure, bandwidth).", None),
        ("Light blue rows 63-68 are Team C additions: margin, SNR at PRFE, e.r.p. and the CNAF check.", None),
        ("Row 43 keeps the course's definition: sensitivity (G31) over noise power, not PRFE over noise.", None),
        ("", None),
        ("Full multi-scenario model: MTP-F26_TeamC_Radio_LinkBudget.xlsx (same maths, 15 scenarios).", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = readme.cell(row=i, column=1, value=text)
        if font:
            c.font = font
    readme.column_dimensions["A"].width = 110
    for title, inputs, note in SHEETS:
        build_sheet(wb.create_sheet(title), inputs, note)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
